"""MeetWise AI — Premium AI Meeting Notes Generator.

Converts lengthy meeting notes into structured summaries: key points,
decisions, action items, deadlines. Powered by Groq.

Run with: streamlit run app.py
"""
from __future__ import annotations

import html
import io
import json
import os
import re
import uuid
from datetime import datetime
from typing import Any, Optional

import pandas as pd
import streamlit as st
from openai import OpenAI

# Optional PDF/DOCX support
try:
    import pymupdf
    PDF_OK = True
except Exception:
    PDF_OK = False

try:
    from docx import Document as DocxDocument
    DOCX_OK = True
except Exception:
    DOCX_OK = False

APP_NAME = "MeetWise AI"
GROQ_BASE_URL = "https://api.groq.com/openai/v1"

DEFAULT_MODEL = "llama-3.3-70b-versatile"
MODEL_FALLBACKS = (
    "llama-3.3-70b-versatile",
    "llama-3.1-8b-instant",
    "openai/gpt-oss-120b",
    "openai/gpt-oss-20b",
    "qwen/qwen3-32b",
    "moonshotai/kimi-k2-instruct",
)

MAX_NOTES_CHARS = 40000
MAX_MEETINGS = 20

# ---------------------------------------------------------------------------
# Meeting templates
# ---------------------------------------------------------------------------
TEMPLATES = {
    "auto": {
        "label": "🔍 Auto-detect",
        "hint": "AI figures out the meeting type from the content.",
        "focus": "",
    },
    "standup": {
        "label": "⚡ Daily Standup",
        "hint": "Short team sync — yesterday, today, blockers.",
        "focus": (
            "This is a short daily standup. Focus on: what each person did "
            "yesterday, what they're doing today, and any blockers."
        ),
    },
    "brainstorm": {
        "label": "💡 Brainstorm",
        "hint": "Idea generation session with many possibilities.",
        "focus": (
            "This is a brainstorming session. Focus on: ideas generated, "
            "pros/cons discussed, top candidates, next experiments."
        ),
    },
    "client": {
        "label": "🤝 Client Call",
        "hint": "External meeting with a client or customer.",
        "focus": (
            "This is an external client/customer call. Focus on: client "
            "needs, commitments we made, follow-up deliverables, pricing/timeline."
        ),
    },
    "one_on_one": {
        "label": "👥 1-on-1",
        "hint": "Manager-report private meeting.",
        "focus": (
            "This is a 1-on-1. Focus on: personal feedback, growth goals, "
            "well-being, career topics, and action items for both parties."
        ),
    },
    "review": {
        "label": "📊 Review / Retro",
        "hint": "Performance review or sprint retrospective.",
        "focus": (
            "This is a review or retrospective. Focus on: what went well, "
            "what didn't, root causes, and specific improvements for next cycle."
        ),
    },
    "board": {
        "label": "🏛 Board / Exec",
        "hint": "Executive or board-level strategic meeting.",
        "focus": (
            "This is a board or executive meeting. Focus on: strategic decisions, "
            "financials, risks, governance, and major approvals."
        ),
    },
    "training": {
        "label": "🎓 Training / Workshop",
        "hint": "Teaching or workshop-style session.",
        "focus": (
            "This is a training or workshop. Focus on: learning objectives, "
            "key concepts taught, exercises, and follow-up practice."
        ),
    },
}

# ---------------------------------------------------------------------------
# Page config
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title=APP_NAME,
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ===========================================================================
# PREMIUM CSS — Indigo/Slate theme
# ===========================================================================
PREMIUM_CSS = """
<style>
:root {
  --mw-ink: #0f172a;
  --mw-ink-2: #475569;
  --mw-ink-3: #94a3b8;
  --mw-bg: #f8fafc;
  --mw-surface: #ffffff;
  --mw-surface-2: #f1f5f9;
  --mw-border: #e2e8f0;
  --mw-border-2: #cbd5e1;
  --mw-primary: #4f46e5;
  --mw-primary-2: #4338ca;
  --mw-primary-3: #a5b4fc;
  --mw-soft: #eef2ff;
  --mw-soft-2: #e0e7ff;
  --mw-accent: #ec4899;
  --mw-accent-2: #db2777;
  --mw-success: #10b981;
  --mw-warning: #f59e0b;
  --mw-danger: #dc2626;
  --mw-info: #06b6d4;
  --mw-shadow-sm: 0 1px 2px rgba(15,23,42,.05);
  --mw-shadow-md: 0 4px 14px rgba(15,23,42,.08);
  --mw-shadow-lg: 0 12px 36px rgba(15,23,42,.10);
  --mw-radius-sm: 8px;
  --mw-radius-md: 12px;
  --mw-radius-lg: 18px;
  --mw-radius-xl: 26px;
}

[data-testid="stAppViewContainer"] { background: var(--mw-bg); }
[data-testid="stHeader"] {
  background: rgba(248,250,252,.8);
  backdrop-filter: blur(8px);
  border-bottom: 1px solid var(--mw-border);
}
.block-container { padding: 1.5rem 2.5rem 4rem; max-width: 1500px; }

/* ---------- Universal readability ---------- */
html, body, .stApp, [data-testid="stAppViewContainer"] {
  color: var(--mw-ink);
}
[data-testid="stAppViewContainer"] p,
[data-testid="stAppViewContainer"] li,
[data-testid="stAppViewContainer"] span,
[data-testid="stAppViewContainer"] label,
[data-testid="stAppViewContainer"] h1,
[data-testid="stAppViewContainer"] h2,
[data-testid="stAppViewContainer"] h3,
[data-testid="stAppViewContainer"] h4,
[data-testid="stAppViewContainer"] h5,
[data-testid="stAppViewContainer"] h6,
[data-testid="stAppViewContainer"] [data-testid="stMarkdownContainer"],
[data-testid="stAppViewContainer"] [data-testid="stMarkdownContainer"] * {
  color: var(--mw-ink);
}
[data-testid="stCaptionContainer"] *,
[data-testid="stAppViewContainer"] small {
  color: var(--mw-ink-2) !important;
}

/* ---------- Inputs ---------- */
.stTextInput input, .stTextArea textarea, .stNumberInput input,
.stChatInput textarea, [data-testid="stChatInput"] textarea,
[data-testid="stTextInput"] input, [data-testid="stTextArea"] textarea,
[data-baseweb="input"] input, [data-baseweb="textarea"] textarea,
[data-baseweb="select"] > div, .stSelectbox div[data-baseweb="select"] > div {
  background-color: #ffffff !important;
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
  caret-color: var(--mw-primary) !important;
  border: 1px solid var(--mw-border-2) !important;
  border-radius: var(--mw-radius-sm) !important;
}
.stTextInput input:focus, .stTextArea textarea:focus {
  border-color: var(--mw-primary) !important;
  box-shadow: 0 0 0 3px var(--mw-soft-2) !important;
}

/* Dropdown popovers */
[data-baseweb="popover"], [role="listbox"], [role="option"] {
  background-color: #ffffff !important;
  color: var(--mw-ink) !important;
}
[role="option"]:hover, [aria-selected="true"] {
  background-color: var(--mw-soft) !important;
  color: var(--mw-primary) !important;
}

/* ---------- Buttons ---------- */
.stButton > button, .stDownloadButton > button {
  background: #ffffff !important;
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
  border: 1px solid var(--mw-border-2) !important;
  border-radius: var(--mw-radius-sm) !important;
  font-weight: 600 !important;
  padding: 0.5rem 1rem !important;
  transition: all .15s ease !important;
  box-shadow: var(--mw-shadow-sm) !important;
}
.stButton > button:hover, .stDownloadButton > button:hover {
  border-color: var(--mw-primary) !important;
  color: var(--mw-primary) !important;
  -webkit-text-fill-color: var(--mw-primary) !important;
  transform: translateY(-1px) !important;
  box-shadow: var(--mw-shadow-md) !important;
}
.stButton > button[kind="primary"],
.stDownloadButton > button[kind="primary"] {
  background: var(--mw-primary) !important;
  color: #ffffff !important;
  -webkit-text-fill-color: #ffffff !important;
  border-color: var(--mw-primary) !important;
}
.stButton > button[kind="primary"]:hover,
.stDownloadButton > button[kind="primary"]:hover {
  background: var(--mw-primary-2) !important;
  border-color: var(--mw-primary-2) !important;
}

/* ---------- Sidebar ---------- */
section[data-testid="stSidebar"] {
  background: linear-gradient(180deg, #1e1b4b 0%, #312e81 100%) !important;
  border-right: 1px solid rgba(255,255,255,.08) !important;
}
section[data-testid="stSidebar"] * {
  color: #e0e7ff !important;
  -webkit-text-fill-color: #e0e7ff !important;
}
section[data-testid="stSidebar"] h1,
section[data-testid="stSidebar"] h2,
section[data-testid="stSidebar"] h3,
section[data-testid="stSidebar"] h4,
section[data-testid="stSidebar"] strong {
  color: #ffffff !important;
  -webkit-text-fill-color: #ffffff !important;
  font-weight: 700 !important;
}
section[data-testid="stSidebar"] small,
section[data-testid="stSidebar"] [data-testid="stCaptionContainer"] * {
  color: #a5b4fc !important;
  -webkit-text-fill-color: #a5b4fc !important;
}
section[data-testid="stSidebar"] hr {
  border-color: rgba(255,255,255,.12) !important;
  margin: 1rem 0 !important;
}
section[data-testid="stSidebar"] .stButton > button {
  background: rgba(255,255,255,.08) !important;
  color: #e0e7ff !important;
  -webkit-text-fill-color: #e0e7ff !important;
  border: 1px solid rgba(255,255,255,.15) !important;
}
section[data-testid="stSidebar"] .stButton > button:hover {
  background: rgba(255,255,255,.16) !important;
  color: #ffffff !important;
  -webkit-text-fill-color: #ffffff !important;
  border-color: var(--mw-primary-3) !important;
}
section[data-testid="stSidebar"] .stButton > button[kind="primary"] {
  background: var(--mw-primary) !important;
  color: #ffffff !important;
  -webkit-text-fill-color: #ffffff !important;
  border-color: var(--mw-primary) !important;
}
section[data-testid="stSidebar"] input,
section[data-testid="stSidebar"] textarea,
section[data-testid="stSidebar"] [data-baseweb="select"] > div {
  background: rgba(255,255,255,.08) !important;
  color: #ffffff !important;
  -webkit-text-fill-color: #ffffff !important;
  border: 1px solid rgba(255,255,255,.15) !important;
}
section[data-testid="stSidebar"] [data-baseweb="popover"],
section[data-testid="stSidebar"] [role="option"] {
  background: #312e81 !important;
  color: #e0e7ff !important;
  -webkit-text-fill-color: #e0e7ff !important;
}
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] {
  background-color: rgba(255,255,255,.06) !important;
  border: 2px dashed rgba(255,255,255,.2) !important;
  border-radius: var(--mw-radius-md) !important;
}
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] *,
section[data-testid="stSidebar"] [data-testid="stFileUploaderFileName"] {
  color: #e0e7ff !important;
  -webkit-text-fill-color: #e0e7ff !important;
}
section[data-testid="stSidebar"] [data-testid="stFileUploaderDropzoneInstructions"] * {
  color: #a5b4fc !important;
  -webkit-text-fill-color: #a5b4fc !important;
}
section[data-testid="stSidebar"] [data-testid="stFileUploaderFile"] {
  background-color: rgba(255,255,255,.06) !important;
  border: 1px solid rgba(255,255,255,.10) !important;
  border-radius: 8px !important;
}
section[data-testid="stSidebar"] [role="slider"] {
  background: var(--mw-primary-3) !important;
}

/* ---------- Metrics ---------- */
[data-testid="stMetric"] {
  background: var(--mw-surface) !important;
  border: 1px solid var(--mw-border) !important;
  border-radius: var(--mw-radius-md) !important;
  padding: 1rem 1.15rem !important;
  box-shadow: var(--mw-shadow-sm) !important;
}
[data-testid="stMetricLabel"] * {
  color: var(--mw-ink-2) !important;
  font-weight: 600 !important;
  text-transform: uppercase !important;
  letter-spacing: 0.05em !important;
  font-size: 0.72rem !important;
}
[data-testid="stMetricValue"] * {
  color: var(--mw-ink) !important;
  font-weight: 750 !important;
  font-size: 1.5rem !important;
}

/* ---------- Tabs ---------- */
[data-testid="stTabs"] [role="tablist"] {
  border-bottom: 1px solid var(--mw-border) !important;
  gap: 0.25rem !important;
}
[data-testid="stTabs"] button[role="tab"] {
  color: var(--mw-ink-2) !important;
  -webkit-text-fill-color: var(--mw-ink-2) !important;
  font-weight: 600 !important;
  border-radius: var(--mw-radius-sm) var(--mw-radius-sm) 0 0 !important;
  padding: 0.6rem 1rem !important;
  border: none !important;
  background: transparent !important;
}
[data-testid="stTabs"] button[role="tab"]:hover {
  color: var(--mw-primary) !important;
  -webkit-text-fill-color: var(--mw-primary) !important;
  background: var(--mw-soft) !important;
}
[data-testid="stTabs"] button[role="tab"][aria-selected="true"] {
  color: var(--mw-primary) !important;
  -webkit-text-fill-color: var(--mw-primary) !important;
  background: var(--mw-soft) !important;
  border-bottom: 2px solid var(--mw-primary) !important;
}

/* ---------- Expander ---------- */
[data-testid="stExpander"] {
  border: 1px solid var(--mw-border) !important;
  border-radius: var(--mw-radius-md) !important;
  background: var(--mw-surface) !important;
  box-shadow: var(--mw-shadow-sm) !important;
  overflow: hidden !important;
}
[data-testid="stExpander"] summary,
[data-testid="stExpander"] summary * {
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
  font-weight: 600 !important;
}

/* ---------- Alerts ---------- */
[data-testid="stAlert"] {
  background: var(--mw-surface) !important;
  border-radius: var(--mw-radius-md) !important;
  border: 1px solid var(--mw-border) !important;
}
[data-testid="stAlert"] * { color: var(--mw-ink) !important; }

/* ---------- Code ---------- */
pre, code, [data-testid="stCode"] * {
  background: var(--mw-surface-2) !important;
  color: var(--mw-ink) !important;
  border-radius: var(--mw-radius-sm) !important;
}

/* ---------- Dataframe ---------- */
[data-testid="stDataFrame"] * {
  color: var(--mw-ink) !important;
  background: #ffffff !important;
}

/* ---------- Custom hero ---------- */
.mw-hero {
  position: relative;
  overflow: hidden;
  padding: 2.25rem 2.5rem;
  margin-bottom: 1.5rem;
  border-radius: var(--mw-radius-xl);
  background:
    radial-gradient(circle at top right, rgba(79,70,229,.15), transparent 50%),
    radial-gradient(circle at bottom left, rgba(236,72,153,.10), transparent 50%),
    linear-gradient(135deg, #ffffff 0%, #f5f3ff 100%);
  border: 1px solid var(--mw-border);
  box-shadow: var(--mw-shadow-lg);
}
.mw-hero::before {
  content: "";
  position: absolute;
  top: 0; left: 0; right: 0;
  height: 4px;
  background: linear-gradient(90deg, #4f46e5, #a5b4fc, #ec4899);
}
.mw-kicker {
  display: inline-block;
  font-size: 0.7rem;
  font-weight: 800;
  letter-spacing: 0.16em;
  text-transform: uppercase;
  color: var(--mw-primary) !important;
  -webkit-text-fill-color: var(--mw-primary) !important;
  background: var(--mw-soft-2);
  padding: 0.4rem 0.85rem;
  border-radius: 999px;
  border: 1px solid #c7d2fe;
  margin-bottom: 1rem;
}
.mw-title {
  font-size: clamp(1.9rem, 3.4vw, 2.75rem);
  font-weight: 800;
  letter-spacing: -0.03em;
  line-height: 1.1;
  margin: 0 0 0.65rem;
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
}
.mw-title .accent {
  background: linear-gradient(90deg, #4f46e5 0%, #8b5cf6 50%, #ec4899 100%);
  -webkit-background-clip: text;
  background-clip: text;
  -webkit-text-fill-color: transparent !important;
  color: transparent !important;
}
.mw-subtitle {
  font-size: 1.05rem;
  line-height: 1.65;
  color: var(--mw-ink-2) !important;
  -webkit-text-fill-color: var(--mw-ink-2) !important;
  max-width: 800px;
  margin: 0;
}

/* ---------- Summary cards ---------- */
.mw-card {
  background: var(--mw-surface);
  border: 1px solid var(--mw-border);
  border-radius: var(--mw-radius-lg);
  padding: 1.5rem 1.6rem;
  margin-bottom: 1.25rem;
  box-shadow: var(--mw-shadow-sm);
  transition: transform .2s ease, box-shadow .2s ease;
}
.mw-card:hover { box-shadow: var(--mw-shadow-md); }
.mw-card.accent-purple { border-top: 3px solid var(--mw-primary); }
.mw-card.accent-pink   { border-top: 3px solid var(--mw-accent); }
.mw-card.accent-green  { border-top: 3px solid var(--mw-success); }
.mw-card.accent-amber  { border-top: 3px solid var(--mw-warning); }
.mw-card.accent-red    { border-top: 3px solid var(--mw-danger); }
.mw-card.accent-cyan   { border-top: 3px solid var(--mw-info); }

.mw-card-head {
  display: flex;
  align-items: center;
  gap: 0.75rem;
  margin-bottom: 0.85rem;
}
.mw-card-icon {
  display: inline-flex;
  align-items: center;
  justify-content: center;
  width: 36px; height: 36px;
  border-radius: 10px;
  background: var(--mw-soft);
  font-size: 1.1rem;
}
.mw-card-title {
  font-size: 0.72rem;
  font-weight: 800;
  letter-spacing: 0.1em;
  text-transform: uppercase;
  color: var(--mw-primary) !important;
  -webkit-text-fill-color: var(--mw-primary) !important;
  margin: 0;
}
.mw-card-body {
  font-size: 1rem;
  line-height: 1.65;
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
}

/* ---------- Action items table ---------- */
.mw-action-row {
  display: grid;
  grid-template-columns: 32px 1.2fr 2.5fr 0.9fr 0.9fr;
  gap: 0.5rem;
  padding: 0.75rem 0.85rem;
  border-bottom: 1px solid var(--mw-border);
  align-items: center;
  font-size: 0.9rem;
}
.mw-action-row:last-child { border-bottom: none; }
.mw-action-row.head {
  background: var(--mw-surface-2);
  font-weight: 700;
  font-size: 0.78rem;
  text-transform: uppercase;
  letter-spacing: 0.06em;
  color: var(--mw-ink-2) !important;
  -webkit-text-fill-color: var(--mw-ink-2) !important;
  border-radius: var(--mw-radius-sm) var(--mw-radius-sm) 0 0;
}
.mw-action-row .prio {
  display: inline-block;
  padding: 0.15rem 0.55rem;
  border-radius: 999px;
  font-size: 0.72rem;
  font-weight: 700;
}
.prio-high   { background: #fee2e2; color: #991b1b !important; -webkit-text-fill-color: #991b1b !important; }
.prio-medium { background: #fef3c7; color: #854d0e !important; -webkit-text-fill-color: #854d0e !important; }
.prio-low    { background: #dcfce7; color: #166534 !important; -webkit-text-fill-color: #166534 !important; }
.prio-na     { background: #f1f5f9; color: #475569 !important; -webkit-text-fill-color: #475569 !important; }

/* ---------- Deadline chips ---------- */
.mw-deadline {
  display: flex;
  align-items: center;
  gap: 0.85rem;
  padding: 0.85rem 1rem;
  border-radius: var(--mw-radius-md);
  background: var(--mw-surface-2);
  border: 1px solid var(--mw-border);
  margin-bottom: 0.55rem;
}
.mw-deadline.urgent { background: #fef2f2; border-color: #fecaca; }
.mw-deadline.soon   { background: #fffbeb; border-color: #fde68a; }
.mw-deadline.later  { background: #f0fdf4; border-color: #bbf7d0; }
.mw-deadline-date {
  min-width: 110px;
  font-weight: 700;
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
  font-size: 0.88rem;
}
.mw-deadline-task {
  flex: 1;
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
  font-size: 0.92rem;
}
.mw-deadline-owner {
  font-size: 0.82rem;
  color: var(--mw-ink-2) !important;
  -webkit-text-fill-color: var(--mw-ink-2) !important;
  font-style: italic;
}

/* ---------- Attendee chips ---------- */
.mw-attendees {
  display: flex;
  flex-wrap: wrap;
  gap: 0.5rem;
}
.mw-attendee {
  display: inline-flex;
  align-items: center;
  gap: 0.4rem;
  background: var(--mw-soft);
  border: 1px solid var(--mw-soft-2);
  border-radius: 999px;
  padding: 0.35rem 0.85rem;
  font-size: 0.85rem;
  font-weight: 600;
  color: var(--mw-primary) !important;
  -webkit-text-fill-color: var(--mw-primary) !important;
}

/* ---------- Decision list ---------- */
.mw-decision {
  padding: 0.85rem 1rem;
  background: #f0fdf4;
  border-left: 3px solid var(--mw-success);
  border-radius: var(--mw-radius-sm);
  margin-bottom: 0.5rem;
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
  font-size: 0.92rem;
  line-height: 1.55;
}

/* ---------- Topic list ---------- */
.mw-topic {
  padding: 0.85rem 1rem;
  background: #f5f3ff;
  border-left: 3px solid var(--mw-primary);
  border-radius: var(--mw-radius-sm);
  margin-bottom: 0.5rem;
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
  font-size: 0.92rem;
  line-height: 1.55;
}

/* ---------- Info card (landing) ---------- */
.mw-info-card {
  padding: 1.35rem 1.4rem;
  border-radius: var(--mw-radius-lg);
  background: var(--mw-surface);
  border: 1px solid var(--mw-border);
  box-shadow: var(--mw-shadow-sm);
  height: 100%;
  transition: transform .2s ease, box-shadow .2s ease;
}
.mw-info-card:hover {
  transform: translateY(-3px);
  box-shadow: var(--mw-shadow-md);
}
.mw-info-icon {
  display: inline-flex;
  align-items: center;
  justify-content: center;
  width: 44px; height: 44px;
  border-radius: 12px;
  background: var(--mw-soft-2);
  color: var(--mw-primary) !important;
  -webkit-text-fill-color: var(--mw-primary) !important;
  font-size: 1.3rem;
  margin-bottom: 0.85rem;
}
.mw-info-card h4 {
  color: var(--mw-ink) !important;
  -webkit-text-fill-color: var(--mw-ink) !important;
  font-size: 1.05rem;
  font-weight: 700;
  margin: 0 0 0.5rem;
}
.mw-info-card p {
  color: var(--mw-ink-2) !important;
  -webkit-text-fill-color: var(--mw-ink-2) !important;
  font-size: 0.92rem;
  line-height: 1.6;
  margin: 0;
}

/* Hide Streamlit branding */
#MainMenu, footer { visibility: hidden; }
[data-testid="stToolbar"] { display: none; }

@media (max-width: 768px) {
  .block-container { padding: 1rem 1rem 3rem !important; }
  .mw-hero { padding: 1.5rem 1.25rem; border-radius: var(--mw-radius-lg); }
  .mw-title { font-size: 1.7rem; }
  .mw-action-row { grid-template-columns: 1fr; }
}
</style>
"""
st.markdown(PREMIUM_CSS, unsafe_allow_html=True)


# ===========================================================================
# API helpers
# ===========================================================================

def resolve_api_key() -> Optional[str]:
    key = os.getenv("GROQ_API_KEY", "").strip()
    if key:
        return key
    try:
        return str(st.secrets.get("GROQ_API_KEY", "")).strip() or None
    except Exception:
        return None


def get_groq_client() -> OpenAI:
    key = resolve_api_key()
    if not key:
        raise RuntimeError(
            "Groq API key is missing. Set GROQ_API_KEY in your environment "
            "or in .streamlit/secrets.toml."
        )
    return OpenAI(api_key=key, base_url=GROQ_BASE_URL)


def _classify_groq_error(exc: Exception) -> tuple[str, str]:
    msg = str(exc).lower()
    if any(w in msg for w in ("401", "unauthorized", "invalid api key", "authentication")):
        return ("auth", "Groq authentication failed. Your GROQ_API_KEY is invalid or expired.")
    if any(w in msg for w in ("402", "insufficient", "payment required")):
        return ("billing", "Groq account has insufficient balance or quota.")
    if any(w in msg for w in ("429", "rate limit", "too many requests")):
        return ("rate_limit", "Groq rate limit reached. Please wait a moment and retry.")
    if any(w in msg for w in ("timeout", "deadline", "connection", "network", "unreachable")):
        return ("network", "Could not reach Groq — network or DNS issue.")
    if any(w in msg for w in ("404", "model_not_found", "does not exist",
                              "decommissioned", "no such model", "invalid model")):
        return ("model", "Model unavailable.")
    if any(w in msg for w in ("context length", "too many tokens", "maximum context", "token limit")):
        return ("context", "Meeting notes are too long for this model. Trim the input or switch models.")
    if any(w in msg for w in ("400", "bad request")):
        return ("bad_request", f"Groq rejected the request: {str(exc)[:200]}")
    return ("unknown", f"Groq request failed: {str(exc)[:200]}")


def test_groq(model: str) -> tuple[bool, str]:
    if not resolve_api_key():
        return False, "❌ GROQ_API_KEY not set."
    try:
        client = get_groq_client()
        resp = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": "Reply with just: OK"}],
            max_tokens=5,
            temperature=0,
        )
        text = (resp.choices[0].message.content or "").strip()
        if not text:
            return False, f"⚠️ Empty response from `{model}`."
        return True, f"✅ `{model}` is working. Response: {text[:30]}"
    except Exception as exc:
        cat, msg = _classify_groq_error(exc)
        return False, f"❌ [{cat}] {msg}"


def call_groq_json(
    prompt: str,
    system_instruction: str,
    model: str,
    max_tokens: int = 3500,
    temperature: float = 0.2,
) -> dict[str, Any]:
    """Call Groq expecting JSON output. Falls back across models."""
    client = get_groq_client()
    candidates = [model] + [m for m in MODEL_FALLBACKS if m != model]
    last_error = "No models tried."

    for model_name in candidates:
        try:
            kwargs: dict[str, Any] = {
                "model": model_name,
                "messages": [
                    {"role": "system", "content": system_instruction.strip()},
                    {"role": "user", "content": prompt},
                ],
                "max_tokens": max_tokens,
                "temperature": temperature,
                "stream": False,
            }
            # Try JSON mode first
            try:
                resp = client.chat.completions.create(
                    **kwargs,
                    response_format={"type": "json_object"},
                )
            except Exception:
                # Retry without JSON mode
                resp = client.chat.completions.create(**kwargs)

            text = (resp.choices[0].message.content or "").strip()
            if not text:
                raise RuntimeError("Empty response.")

            data = _parse_json_safe(text)
            if isinstance(data, dict):
                try:
                    st.session_state["active_model"] = model_name
                    usage = getattr(resp, "usage", None)
                    if usage is not None:
                        st.session_state["last_usage"] = {
                            "prompt": getattr(usage, "prompt_tokens", 0) or 0,
                            "completion": getattr(usage, "completion_tokens", 0) or 0,
                            "total": getattr(usage, "total_tokens", 0) or 0,
                        }
                except Exception:
                    pass
                return data
            raise RuntimeError(f"Model returned non-JSON: {text[:150]}")
        except Exception as exc:
            cat, msg = _classify_groq_error(exc)
            last_error = f"[{cat}] {msg}"
            if cat in ("model", "context"):
                continue
            raise RuntimeError(msg)

    raise RuntimeError(f"All Groq models failed. Last: {last_error}")


def _parse_json_safe(text: str) -> Any:
    if not text:
        return None
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
    cleaned = re.sub(r"\s*```\s*$", "", cleaned)
    try:
        return json.loads(cleaned)
    except Exception:
        pass
    s, e = cleaned.find("{"), cleaned.rfind("}")
    if s != -1 and e != -1 and e > s:
        try:
            return json.loads(cleaned[s:e + 1])
        except Exception:
            pass
    return None


# ===========================================================================
# File extraction
# ===========================================================================

def extract_from_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1", "utf-16"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_from_pdf(data: bytes) -> str:
    if not PDF_OK:
        raise RuntimeError("PyMuPDF not installed. Run: pip install PyMuPDF")
    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        parts = []
        for page in doc:
            text = page.get_text("text")
            if text:
                parts.append(text)
        return "\n\n".join(parts)
    finally:
        doc.close()


def extract_from_docx(data: bytes) -> str:
    if not DOCX_OK:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")
    doc = DocxDocument(io.BytesIO(data))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_from_upload(uploaded_file: Any) -> str:
    """Auto-detect file type and extract text."""
    data = uploaded_file.getvalue()
    name = (getattr(uploaded_file, "name", "") or "").lower()
    if name.endswith(".pdf"):
        return extract_from_pdf(data)
    if name.endswith(".docx"):
        return extract_from_docx(data)
    if name.endswith(".doc"):
        raise RuntimeError("Old .doc format not supported. Save as .docx or .pdf.")
    return extract_from_txt(data)


# ===========================================================================
# Meeting notes extraction
# ===========================================================================

EXTRACTION_SYSTEM = """You are MeetWise AI, an expert meeting-notes analyst.

You convert raw meeting transcripts or notes into a clean, structured JSON summary.

Return ONLY valid JSON (no markdown fences, no commentary) matching this exact schema:

{
  "meeting_title": "short descriptive title (5-12 words)",
  "meeting_date": "YYYY-MM-DD or empty string if unknown",
  "meeting_type": "one of: standup, brainstorm, client, one_on_one, review, board, training, general",
  "duration_minutes": null or integer,
  "attendees": [
    {"name": "Full Name", "role": "role or role hint, if mentioned"}
  ],
  "executive_summary": "2-4 sentence high-level overview of the meeting",
  "key_points": [
    "concise bullet capturing a major discussion point"
  ],
  "topics": [
    {"topic": "topic name", "discussion": "1-2 sentence summary of what was discussed"}
  ],
  "decisions": [
    "specific decision that was made (stated as fact)"
  ],
  "action_items": [
    {
      "task": "clear, imperative description of the task",
      "owner": "person name or 'Unassigned'",
      "due_date": "YYYY-MM-DD or relative phrase like 'next Friday' or empty string",
      "priority": "high|medium|low",
      "context": "brief note on why/what"
    }
  ],
  "deadlines": [
    {"date": "YYYY-MM-DD or text", "what": "what is due", "owner": "person or empty"}
  ],
  "open_questions": [
    "question raised but not resolved"
  ],
  "risks_blockers": [
    "risk, blocker, or concern raised"
  ],
  "next_meeting": "date/time or empty string if not mentioned",
  "sentiment": "positive|neutral|mixed|negative",
  "readability_notes": "1 sentence on how clear the original notes were"
}

Rules:
- NEVER invent attendees, decisions, or action items. If unsure, omit or use empty string.
- Action item tasks MUST start with an imperative verb (e.g., "Send", "Review", "Update").
- Priority: infer from urgency cues ("ASAP" → high, "when possible" → low). Default to medium.
- Due dates: if relative ("next Monday"), keep the phrase as-is AND add an ISO date if you can infer.
- Multiple action items from one sentence should be split into separate entries.
- If the input is short or unclear, still return the full schema with empty arrays where needed.
- Preserve exact names, numbers, and technical terms from the source."""


def analyze_meeting(
    notes: str,
    template_key: str = "auto",
    language: str = "English",
    model: str = DEFAULT_MODEL,
) -> dict[str, Any]:
    if not notes or not notes.strip():
        raise ValueError("Notes are empty.")

    notes = notes.strip()
    if len(notes) > MAX_NOTES_CHARS:
        notes = notes[:MAX_NOTES_CHARS] + "\n\n[truncated]"

    template = TEMPLATES.get(template_key, TEMPLATES["auto"])
    focus_note = f"\n\nFOCUS: {template['focus']}" if template["focus"] else ""

    lang_note = ""
    if language == "Roman Urdu":
        lang_note = (
            "\n\nLANGUAGE: Write all text fields (summaries, tasks, decisions) in "
            "ROMAN URDU (Urdu written in English letters). Keep JSON keys in English."
        )

    prompt = f"""Analyze the following meeting notes and produce the structured JSON.

Focus context: {focus_note or 'Auto-detect the meeting type.'}

MEETING NOTES:
---
{notes}
---

Return the JSON schema as specified. Be faithful to the source — no inventions."""

    system = EXTRACTION_SYSTEM + lang_note
    data = call_groq_json(prompt, system, model=model, max_tokens=4000, temperature=0.2)
    return _normalize_meeting(data, notes)


def _normalize_meeting(data: dict[str, Any], raw_notes: str) -> dict[str, Any]:
    def as_list(x: Any) -> list:
        if isinstance(x, list):
            return x
        if x is None:
            return []
        return [x]

    def as_str(x: Any, default: str = "") -> str:
        if x is None:
            return default
        return str(x).strip()

    out: dict[str, Any] = {
        "meeting_title": as_str(data.get("meeting_title")) or _auto_title(raw_notes),
        "meeting_date": as_str(data.get("meeting_date")),
        "meeting_type": as_str(data.get("meeting_type")) or "general",
        "duration_minutes": data.get("duration_minutes") if isinstance(data.get("duration_minutes"), int) else None,
        "attendees": [],
        "executive_summary": as_str(data.get("executive_summary")),
        "key_points": [as_str(x) for x in as_list(data.get("key_points")) if as_str(x)],
        "topics": [],
        "decisions": [as_str(x) for x in as_list(data.get("decisions")) if as_str(x)],
        "action_items": [],
        "deadlines": [],
        "open_questions": [as_str(x) for x in as_list(data.get("open_questions")) if as_str(x)],
        "risks_blockers": [as_str(x) for x in as_list(data.get("risks_blockers")) if as_str(x)],
        "next_meeting": as_str(data.get("next_meeting")),
        "sentiment": as_str(data.get("sentiment")) or "neutral",
        "readability_notes": as_str(data.get("readability_notes")),
        "raw_notes": raw_notes,
        "extracted_at": datetime.now().isoformat(),
    }

    for a in as_list(data.get("attendees")):
        if isinstance(a, dict):
            name = as_str(a.get("name"))
            if name:
                out["attendees"].append({
                    "name": name,
                    "role": as_str(a.get("role")),
                })
        elif isinstance(a, str) and a.strip():
            out["attendees"].append({"name": a.strip(), "role": ""})

    for t in as_list(data.get("topics")):
        if isinstance(t, dict):
            topic = as_str(t.get("topic"))
            if topic:
                out["topics"].append({
                    "topic": topic,
                    "discussion": as_str(t.get("discussion")),
                })
        elif isinstance(t, str) and t.strip():
            out["topics"].append({"topic": t.strip(), "discussion": ""})

    for a in as_list(data.get("action_items")):
        if isinstance(a, dict):
            task = as_str(a.get("task"))
            if task:
                out["action_items"].append({
                    "task": task,
                    "owner": as_str(a.get("owner")) or "Unassigned",
                    "due_date": as_str(a.get("due_date")),
                    "priority": (as_str(a.get("priority")) or "medium").lower(),
                    "context": as_str(a.get("context")),
                })
        elif isinstance(a, str) and a.strip():
            out["action_items"].append({
                "task": a.strip(), "owner": "Unassigned",
                "due_date": "", "priority": "medium", "context": "",
            })

    for d in as_list(data.get("deadlines")):
        if isinstance(d, dict):
            what = as_str(d.get("what"))
            if what:
                out["deadlines"].append({
                    "date": as_str(d.get("date")),
                    "what": what,
                    "owner": as_str(d.get("owner")),
                })

    return out


def _auto_title(notes: str) -> str:
    first_line = next((ln.strip() for ln in notes.splitlines() if ln.strip()), "Meeting")
    return first_line[:60] + ("…" if len(first_line) > 60 else "")


def enhance_readability(notes: str, model: str = DEFAULT_MODEL) -> str:
    if not notes or len(notes) < 50:
        return notes
    if len(notes) > MAX_NOTES_CHARS:
        notes = notes[:MAX_NOTES_CHARS]

    prompt = f"""Clean up and improve the readability of the following meeting notes.

Rules:
- Fix obvious typos, grammar, and run-on sentences.
- Remove filler words ("um", "uh", "you know", "like" as filler, repetitions).
- Break into clear sections and paragraphs.
- Use bullet points for lists.
- Preserve ALL facts, names, decisions, numbers — do not add or remove information.
- Keep the same language as input (English or Roman Urdu).
- Output ONLY the cleaned text, no commentary, no markdown fences.

NOTES:
---
{notes}
---"""

    system = (
        "You are a professional editor. You improve readability of raw notes "
        "without changing the meaning. Never invent content."
    )

    client = get_groq_client()
    candidates = [model] + [m for m in MODEL_FALLBACKS if m != model]
    for model_name in candidates:
        try:
            resp = client.chat.completions.create(
                model=model_name,
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": prompt},
                ],
                max_tokens=3500,
                temperature=0.3,
                stream=False,
            )
            out = (resp.choices[0].message.content or "").strip()
            if out:
                return out
        except Exception:
            continue
    return notes


# ===========================================================================
# Exports
# ===========================================================================

def build_markdown_report(meeting: dict[str, Any]) -> str:
    m = meeting
    lines = [
        f"# 📋 {m.get('meeting_title', 'Meeting')}",
        "",
        f"**Date:** {m.get('meeting_date') or '—'}  ",
        f"**Type:** {m.get('meeting_type', 'general').title()}  ",
        f"**Sentiment:** {m.get('sentiment', 'neutral').title()}  ",
    ]
    if m.get("duration_minutes"):
        lines.append(f"**Duration:** {m['duration_minutes']} min  ")
    lines.append("")
    lines.append("---")
    lines.append("")

    if m.get("attendees"):
        lines.append("## 👥 Attendees")
        lines.append("")
        for a in m["attendees"]:
            role = f" — *{a['role']}*" if a.get("role") else ""
            lines.append(f"- {a['name']}{role}")
        lines.append("")

    if m.get("executive_summary"):
        lines.append("## 🎯 Executive Summary")
        lines.append("")
        lines.append(m["executive_summary"])
        lines.append("")

    if m.get("key_points"):
        lines.append("## 🔑 Key Discussion Points")
        lines.append("")
        for k in m["key_points"]:
            lines.append(f"- {k}")
        lines.append("")

    if m.get("topics"):
        lines.append("## 📚 Topics Discussed")
        lines.append("")
        for t in m["topics"]:
            lines.append(f"### {t['topic']}")
            lines.append("")
            if t.get("discussion"):
                lines.append(t["discussion"])
                lines.append("")

    if m.get("decisions"):
        lines.append("## ✅ Decisions Made")
        lines.append("")
        for d in m["decisions"]:
            lines.append(f"- {d}")
        lines.append("")

    if m.get("action_items"):
        lines.append("## 📌 Action Items")
        lines.append("")
        lines.append("| # | Task | Owner | Due | Priority |")
        lines.append("|---|------|-------|-----|----------|")
        for i, a in enumerate(m["action_items"], 1):
            task = a["task"].replace("|", "\\|")
            owner = a.get("owner", "Unassigned").replace("|", "\\|")
            due = a.get("due_date", "") or "—"
            prio = a.get("priority", "medium").title()
            lines.append(f"| {i} | {task} | {owner} | {due} | {prio} |")
        lines.append("")

    if m.get("deadlines"):
        lines.append("## ⏰ Deadlines")
        lines.append("")
        for d in m["deadlines"]:
            owner = f" — *{d['owner']}*" if d.get("owner") else ""
            lines.append(f"- **{d.get('date', '—')}** — {d['what']}{owner}")
        lines.append("")

    if m.get("open_questions"):
        lines.append("## ❓ Open Questions")
        lines.append("")
        for q in m["open_questions"]:
            lines.append(f"- {q}")
        lines.append("")

    if m.get("risks_blockers"):
        lines.append("## ⚠️ Risks & Blockers")
        lines.append("")
        for r in m["risks_blockers"]:
            lines.append(f"- {r}")
        lines.append("")

    if m.get("next_meeting"):
        lines.append("## 📅 Next Meeting")
        lines.append("")
        lines.append(m["next_meeting"])
        lines.append("")

    lines.append("---")
    lines.append("")
    lines.append(f"*Generated by MeetWise AI on {datetime.now().strftime('%Y-%m-%d %H:%M')}*")
    return "\n".join(lines)


def build_json_report(meeting: dict[str, Any]) -> str:
    out = {k: v for k, v in meeting.items() if k != "raw_notes"}
    return json.dumps(out, ensure_ascii=False, indent=2)


def build_actions_csv(meeting: dict[str, Any]) -> str:
    rows = []
    for a in meeting.get("action_items", []):
        rows.append({
            "Task": a.get("task", ""),
            "Owner": a.get("owner", "Unassigned"),
            "Due Date": a.get("due_date", ""),
            "Priority": a.get("priority", "medium"),
            "Context": a.get("context", ""),
        })
    if not rows:
        return "Task,Owner,Due Date,Priority,Context\n"
    return pd.DataFrame(rows).to_csv(index=False)


def _safe_pdf(s: Any) -> str:
    """Make text safe for FPDF's latin-1 encoding."""
    return str(s).encode("latin-1", "ignore").decode("latin-1")


def build_pdf_report(meeting: dict[str, Any]) -> bytes:
    """Build a styled PDF report. Returns bytes. Robust against layout errors."""
    try:
        from fpdf import FPDF
    except Exception:
        return b""

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Page geometry
    PAGE_W = 210
    LEFT_MARGIN = 10
    RIGHT_MARGIN = 10
    CONTENT_W = PAGE_W - LEFT_MARGIN - RIGHT_MARGIN  # = 190 mm

    def reset_x():
        pdf.set_x(LEFT_MARGIN)

    def multi(text: str, height: float = 6, font_size: int = 11,
              bold: bool = False, italic: bool = False):
        reset_x()
        style = "B" if bold else ("I" if italic else "")
        pdf.set_font("Helvetica", style, font_size)
        pdf.multi_cell(CONTENT_W, height, _safe_pdf(text),
                       new_x="LMARGIN", new_y="NEXT")
        reset_x()

    def bullet(text: str, height: float = 6, font_size: int = 11):
        reset_x()
        pdf.set_font("Helvetica", "", font_size)
        pdf.multi_cell(CONTENT_W, height, _safe_pdf("- " + text),
                       new_x="LMARGIN", new_y="NEXT")
        reset_x()

    def section(title: str):
        reset_x()
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(79, 70, 229)
        pdf.multi_cell(CONTENT_W, 8, _safe_pdf(title),
                       new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(15, 23, 42)
        reset_x()
        pdf.ln(1)

    def divider():
        reset_x()
        pdf.set_draw_color(226, 232, 240)
        y = pdf.get_y()
        pdf.line(LEFT_MARGIN, y, PAGE_W - RIGHT_MARGIN, y)
        pdf.ln(3)

    # ---------- Header band ----------
    pdf.set_fill_color(79, 70, 229)
    pdf.rect(0, 0, PAGE_W, 30, "F")
    pdf.set_xy(LEFT_MARGIN, 8)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(CONTENT_W, 8, _safe_pdf("MeetWise AI - Meeting Report"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_x(LEFT_MARGIN)
    pdf.set_font("Helvetica", "I", 10)
    pdf.cell(CONTENT_W, 6, _safe_pdf("AI-generated structured summary"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(15, 23, 42)
    reset_x()
    pdf.set_y(38)

    # ---------- Title ----------
    title = meeting.get("meeting_title", "Meeting") or "Meeting"
    multi(title, height=9, font_size=16, bold=True)
    pdf.ln(2)

    # ---------- Meta line ----------
    meta_bits = []
    if meeting.get("meeting_date"):
        meta_bits.append(f"Date: {meeting['meeting_date']}")
    if meeting.get("meeting_type"):
        meta_bits.append(f"Type: {meeting['meeting_type'].title()}")
    if meeting.get("sentiment"):
        meta_bits.append(f"Sentiment: {meeting['sentiment'].title()}")
    if meeting.get("duration_minutes"):
        meta_bits.append(f"Duration: {meeting['duration_minutes']} min")
    if meta_bits:
        multi("  |  ".join(meta_bits), height=6, font_size=10, italic=True)
    pdf.ln(3)
    divider()

    # ---------- Attendees ----------
    if meeting.get("attendees"):
        section("Attendees")
        for a in meeting["attendees"]:
            name = a.get("name", "")
            role = a.get("role", "")
            line = name + (f"  ({role})" if role else "")
            bullet(line)
        pdf.ln(2)

    # ---------- Executive summary ----------
    if meeting.get("executive_summary"):
        section("Executive Summary")
        multi(meeting["executive_summary"])
        pdf.ln(3)

    # ---------- Key points ----------
    if meeting.get("key_points"):
        section("Key Discussion Points")
        for k in meeting["key_points"]:
            bullet(k)
        pdf.ln(2)

    # ---------- Topics ----------
    if meeting.get("topics"):
        section("Topics Discussed")
        for t in meeting["topics"]:
            topic = t.get("topic", "")
            if topic:
                multi(topic, bold=True)
            if t.get("discussion"):
                multi(t["discussion"], font_size=10)
            pdf.ln(1)
        pdf.ln(2)

    # ---------- Decisions ----------
    if meeting.get("decisions"):
        section("Decisions Made")
        for d in meeting["decisions"]:
            bullet(d)
        pdf.ln(2)

    # ---------- Action items ----------
    if meeting.get("action_items"):
        section("Action Items")
        for i, a in enumerate(meeting["action_items"], 1):
            task = a.get("task", "")
            owner = a.get("owner", "Unassigned")
            due = a.get("due_date", "") or "-"
            prio = (a.get("priority", "medium") or "medium").title()
            line = f"{i}. {task}\n   Owner: {owner}  |  Due: {due}  |  Priority: {prio}"
            multi(line, font_size=10)
            pdf.ln(1)
        pdf.ln(2)

    # ---------- Deadlines ----------
    if meeting.get("deadlines"):
        section("Deadlines")
        for d in meeting["deadlines"]:
            date_str = d.get("date", "") or "-"
            what = d.get("what", "")
            owner = d.get("owner", "")
            line = f"{date_str} - {what}"
            if owner:
                line += f"  ({owner})"
            bullet(line)
        pdf.ln(2)

    # ---------- Open questions ----------
    if meeting.get("open_questions"):
        section("Open Questions")
        for q in meeting["open_questions"]:
            bullet(q)
        pdf.ln(2)

    # ---------- Risks ----------
    if meeting.get("risks_blockers"):
        section("Risks & Blockers")
        for r in meeting["risks_blockers"]:
            bullet(r)
        pdf.ln(2)

    # ---------- Next meeting ----------
    if meeting.get("next_meeting"):
        section("Next Meeting")
        multi(meeting["next_meeting"])
        pdf.ln(2)

    # ---------- Footer ----------
    divider()
    multi(
        f"Generated by MeetWise AI on {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        height=5, font_size=9, italic=True,
    )

    out = pdf.output(dest="S")
    if isinstance(out, str):
        return out.encode("latin-1", "ignore")
    return bytes(out)


# ===========================================================================
# Session state
# ===========================================================================

def init_state() -> None:
    defaults: dict[str, Any] = {
        "meetings": {},
        "active_id": None,
        "template": "auto",
        "language": "English",
        "preferred_model": DEFAULT_MODEL,
        "enhance_readability": True,
        "show_debug": False,
        "connection_test": None,
        "active_model": None,
        "last_usage": None,
        "last_error": None,
        "input_text": "",
        "last_upload_name": None,
        "input_version": 0,       # bumped to force text_area refresh
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


def save_meeting(meeting: dict[str, Any]) -> str:
    mid = str(uuid.uuid4())[:8]
    st.session_state.meetings[mid] = meeting
    st.session_state.active_id = mid
    if len(st.session_state.meetings) > MAX_MEETINGS:
        oldest = sorted(
            st.session_state.meetings.items(),
            key=lambda kv: kv[1].get("extracted_at", ""),
        )[:1]
        for oid, _ in oldest:
            del st.session_state.meetings[oid]
    return mid


def get_active_meeting() -> Optional[dict[str, Any]]:
    aid = st.session_state.active_id
    if aid and aid in st.session_state.meetings:
        return st.session_state.meetings[aid]
    return None


# ===========================================================================
# UI Helpers
# ===========================================================================

def esc(s: Any) -> str:
    return html.escape(str(s) if s is not None else "")


def render_hero() -> None:
    st.markdown(
        """
        <div class="mw-hero">
          <div class="mw-kicker">Structured Meeting Intelligence</div>
          <div class="mw-title">📋 MeetWise <span class="accent">AI</span></div>
          <div class="mw-subtitle">
            Turn long, messy meeting notes into clean structured summaries —
            key points, decisions, action items, deadlines, and attendees.
            Powered by Groq.
          </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_landing() -> None:
    st.markdown("### How MeetWise works")
    a, b, c, d = st.columns(4)
    cards = [
        ("📥", "Paste or upload", "Drop in raw notes, a transcript, or a file (.txt, .md, .pdf, .docx)."),
        ("🧠", "AI extraction", "Groq finds key points, decisions, action items, and deadlines."),
        ("🖥️", "Structured output", "Clean cards, tables, and timelines — not a wall of text."),
        ("📤", "Export anywhere", "Markdown, JSON, CSV of action items, or a styled PDF report."),
    ]
    for col, (icon, title, desc) in zip([a, b, c, d], cards):
        with col:
            st.markdown(
                f'<div class="mw-info-card">'
                f'<div class="mw-info-icon">{icon}</div>'
                f'<h4>{title}</h4>'
                f'<p>{desc}</p>'
                f'</div>',
                unsafe_allow_html=True,
            )
    st.markdown("")
    st.info(
        "💡 **Tip:** Works best with notes that have speaker names, timestamps, "
        "or clear action cues ('Alice will send…', 'We decided…', 'Due by Friday')."
    )


def render_meeting(meeting: dict[str, Any]) -> None:
    m = meeting

    st.markdown(f"## {esc(m.get('meeting_title', 'Meeting'))}")
    meta_cols = st.columns(4)
    with meta_cols[0]:
        st.metric("Type", (m.get("meeting_type") or "general").title())
    with meta_cols[1]:
        st.metric("Date", m.get("meeting_date") or "—")
    with meta_cols[2]:
        st.metric("Attendees", len(m.get("attendees", [])))
    with meta_cols[3]:
        st.metric("Action Items", len(m.get("action_items", [])))

    if m.get("executive_summary"):
        st.markdown(
            f'<div class="mw-card accent-purple">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">🎯</div>'
            f'<p class="mw-card-title">Executive Summary</p>'
            f'</div>'
            f'<div class="mw-card-body">{esc(m["executive_summary"])}</div>'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("attendees"):
        chips = "".join(
            f'<span class="mw-attendee">👤 {esc(a["name"])}'
            + (f' <span style="opacity:.7;font-weight:500;">· {esc(a["role"])}</span>' if a.get("role") else "")
            + '</span>'
            for a in m["attendees"]
        )
        st.markdown(
            f'<div class="mw-card accent-cyan">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">👥</div>'
            f'<p class="mw-card-title">Attendees</p>'
            f'</div>'
            f'<div class="mw-attendees">{chips}</div>'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("key_points"):
        items = "".join(f"<li style='margin-bottom:0.4rem;'>{esc(k)}</li>" for k in m["key_points"])
        st.markdown(
            f'<div class="mw-card accent-purple">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">🔑</div>'
            f'<p class="mw-card-title">Key Discussion Points</p>'
            f'</div>'
            f'<div class="mw-card-body"><ul style="margin:0;padding-left:1.25rem;">{items}</ul></div>'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("topics"):
        blocks = ""
        for t in m["topics"]:
            blocks += (
                f'<div class="mw-topic">'
                f'<strong>{esc(t["topic"])}</strong>'
                + (f'<br>{esc(t["discussion"])}' if t.get("discussion") else "")
                + '</div>'
            )
        st.markdown(
            f'<div class="mw-card accent-pink">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">📚</div>'
            f'<p class="mw-card-title">Topics Discussed</p>'
            f'</div>'
            f'{blocks}'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("decisions"):
        items = "".join(f'<div class="mw-decision">✅ {esc(d)}</div>' for d in m["decisions"])
        st.markdown(
            f'<div class="mw-card accent-green">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">✅</div>'
            f'<p class="mw-card-title">Decisions Made</p>'
            f'</div>'
            f'{items}'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("action_items"):
        rows = (
            '<div class="mw-action-row head">'
            '<span>#</span><span>Owner</span><span>Task</span><span>Due</span><span>Priority</span>'
            '</div>'
        )
        for i, a in enumerate(m["action_items"], 1):
            prio = (a.get("priority") or "medium").lower()
            if prio not in ("high", "medium", "low"):
                prio = "na"
            rows += (
                f'<div class="mw-action-row">'
                f'<span style="color:#94a3b8;font-weight:700;">{i}</span>'
                f'<span><strong>{esc(a.get("owner","Unassigned"))}</strong></span>'
                f'<span>{esc(a.get("task",""))}'
                + (f'<br><span style="font-size:0.78rem;color:#94a3b8;">{esc(a.get("context",""))}</span>' if a.get("context") else "")
                + '</span>'
                f'<span>{esc(a.get("due_date") or "—")}</span>'
                f'<span><span class="prio prio-{prio}">{prio.title()}</span></span>'
                f'</div>'
            )
        st.markdown(
            f'<div class="mw-card accent-amber">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">📌</div>'
            f'<p class="mw-card-title">Action Items</p>'
            f'</div>'
            f'<div>{rows}</div>'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("deadlines"):
        items = ""
        for d in m["deadlines"]:
            urgency = "later"
            date_str = (d.get("date") or "").strip()
            iso = re.search(r"\d{4}-\d{2}-\d{2}", date_str)
            if iso:
                try:
                    dt = datetime.strptime(iso.group(0), "%Y-%m-%d")
                    delta = (dt - datetime.now()).days
                    if delta < 0:
                        urgency = "urgent"
                    elif delta <= 7:
                        urgency = "soon"
                except Exception:
                    pass
            else:
                low = date_str.lower()
                if any(w in low for w in ("asap", "today", "tomorrow", "urgent")):
                    urgency = "urgent"
                elif any(w in low for w in ("this week", "next week")):
                    urgency = "soon"
            items += (
                f'<div class="mw-deadline {urgency}">'
                f'<span class="mw-deadline-date">{esc(date_str or "—")}</span>'
                f'<span class="mw-deadline-task">{esc(d.get("what",""))}</span>'
                + (f'<span class="mw-deadline-owner">{esc(d.get("owner",""))}</span>' if d.get("owner") else "")
                + '</div>'
            )
        st.markdown(
            f'<div class="mw-card accent-red">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">⏰</div>'
            f'<p class="mw-card-title">Deadlines & Timeline</p>'
            f'</div>'
            f'{items}'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("open_questions"):
        items = "".join(f"<li>{esc(q)}</li>" for q in m["open_questions"])
        st.markdown(
            f'<div class="mw-card accent-cyan">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">❓</div>'
            f'<p class="mw-card-title">Open Questions</p>'
            f'</div>'
            f'<div class="mw-card-body"><ul style="margin:0;padding-left:1.25rem;">{items}</ul></div>'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("risks_blockers"):
        items = "".join(f"<li>{esc(r)}</li>" for r in m["risks_blockers"])
        st.markdown(
            f'<div class="mw-card accent-red">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">⚠️</div>'
            f'<p class="mw-card-title">Risks & Blockers</p>'
            f'</div>'
            f'<div class="mw-card-body"><ul style="margin:0;padding-left:1.25rem;">{items}</ul></div>'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("next_meeting"):
        st.markdown(
            f'<div class="mw-card accent-purple">'
            f'<div class="mw-card-head">'
            f'<div class="mw-card-icon">📅</div>'
            f'<p class="mw-card-title">Next Meeting</p>'
            f'</div>'
            f'<div class="mw-card-body">{esc(m["next_meeting"])}</div>'
            f'</div>',
            unsafe_allow_html=True,
        )

    if m.get("enhanced_notes"):
        with st.expander("📝 Enhanced notes (readability-improved)"):
            st.markdown(m["enhanced_notes"])

    if st.session_state.get("show_debug"):
        with st.expander("🐛 Raw extracted JSON"):
            st.json({k: v for k, v in m.items() if k != "raw_notes"})
        with st.expander("🐛 Original notes"):
            st.text_area("Raw input", m.get("raw_notes", ""), height=200, disabled=True)


# ===========================================================================
# Main
# ===========================================================================

def main() -> None:
    init_state()
    render_hero()

    # ---------------- Sidebar ----------------
    with st.sidebar:
        st.markdown("### 📥 Input")
        input_mode = st.radio(
            "Source",
            ["✍️ Paste notes", "📎 Upload file"],
            index=0,
            label_visibility="collapsed",
            key="input_mode_radio",
        )

        uploaded_file = None
        if input_mode == "📎 Upload file":
            uploaded_file = st.file_uploader(
                "Upload notes",
                type=["txt", "md", "pdf", "docx"],
                label_visibility="collapsed",
                key="notes_uploader",
            )
            if uploaded_file:
                # Process only when a new file is uploaded
                if st.session_state.last_upload_name != uploaded_file.name:
                    try:
                        text = extract_from_upload(uploaded_file)
                        if not text or not text.strip():
                            st.warning("⚠️ Could not extract any text from this file.")
                        else:
                            st.session_state.input_text = text
                            st.session_state.last_upload_name = uploaded_file.name
                            st.session_state.input_version += 1
                            st.success(f"✅ Loaded {len(text):,} chars from {uploaded_file.name}")
                    except Exception as e:
                        st.error(f"❌ Could not read file: {e}")
            else:
                # No file — reset the tracker so next upload triggers again
                if st.session_state.last_upload_name:
                    st.session_state.last_upload_name = None

        st.markdown("---")
        st.markdown("### 🎯 Meeting Type")
        template_labels = {k: v["label"] for k, v in TEMPLATES.items()}
        current_t = st.session_state.get("template", "auto")
        t_label = st.selectbox(
            "Template",
            options=list(template_labels.values()),
            index=list(template_labels.keys()).index(current_t) if current_t in template_labels else 0,
            label_visibility="collapsed",
        )
        st.session_state.template = [k for k, v in template_labels.items() if v == t_label][0]
        st.caption(TEMPLATES[st.session_state.template]["hint"])

        st.markdown("---")
        st.markdown("### 🌐 Language")
        st.selectbox(
            "Output language",
            options=["English", "Roman Urdu"],
            index=0,
            key="language",
        )

        st.markdown("### ⚙️ Model")
        st.selectbox(
            "Model",
            options=list(MODEL_FALLBACKS),
            index=0,
            key="preferred_model",
        )

        st.checkbox("Enhance readability", key="enhance_readability",
                    help="Also generate a readability-improved version of the notes.")

        st.markdown("---")
        st.markdown("### 📚 History")
        if st.session_state.meetings:
            meetings_sorted = sorted(
                st.session_state.meetings.items(),
                key=lambda kv: kv[1].get("extracted_at", ""),
                reverse=True,
            )[:MAX_MEETINGS]
            for mid, mtg in meetings_sorted:
                is_active = mid == st.session_state.active_id
                col_a, col_b = st.columns([4, 1])
                with col_a:
                    title = mtg.get("meeting_title", "Meeting")[:28]
                    if st.button(
                        title,
                        key=f"mtg_{mid}",
                        use_container_width=True,
                        type="primary" if is_active else "secondary",
                    ):
                        st.session_state.active_id = mid
                        st.rerun()
                with col_b:
                    if st.button("✕", key=f"delmtg_{mid}", use_container_width=True):
                        del st.session_state.meetings[mid]
                        if st.session_state.active_id == mid:
                            remaining = list(st.session_state.meetings.keys())
                            st.session_state.active_id = remaining[0] if remaining else None
                        st.rerun()
            if st.button("🗑️ Clear all history", use_container_width=True):
                st.session_state.meetings = {}
                st.session_state.active_id = None
                st.rerun()
        else:
            st.caption("No meetings yet.")

        st.markdown("---")
        st.markdown("### 🔌 API Status")
        if resolve_api_key():
            st.success("Groq API key detected")
            col_t, col_c = st.columns(2)
            with col_t:
                if st.button("Test", use_container_width=True):
                    with st.spinner("Testing…"):
                        ok, msg = test_groq(st.session_state.get("preferred_model", DEFAULT_MODEL))
                        st.session_state.connection_test = (ok, msg)
            with col_c:
                if st.button("Clear", use_container_width=True):
                    st.session_state.connection_test = None
            if st.session_state.get("connection_test"):
                ok, msg = st.session_state["connection_test"]
                (st.success if ok else st.error)(msg)
            if st.session_state.get("active_model"):
                st.caption(f"Last used: `{st.session_state['active_model']}`")
            usage = st.session_state.get("last_usage")
            if usage:
                st.caption(f"Tokens — prompt: {usage.get('prompt',0)} · completion: {usage.get('completion',0)}")
        else:
            st.warning("Groq API key not configured")

        st.markdown("---")
        st.markdown("### 🐛 Debug")
        st.checkbox("Show debug info", key="show_debug")
        if st.session_state.get("last_error") and st.session_state.get("show_debug"):
            with st.expander("⚠️ Last error", expanded=True):
                st.code(st.session_state["last_error"])

    # ---------------- Main ----------------
    tab_input, tab_result, tab_export = st.tabs(
        ["📝 Input", "📋 Structured Summary", "📤 Export"]
    )

    # ---------- Tab: Input ----------
    with tab_input:
        st.subheader("Meeting notes")
        st.caption(
            "Paste your raw notes/transcript below, or upload a file from the sidebar. "
            f"Max {MAX_NOTES_CHARS:,} characters."
        )

        # Version-controlled key forces Streamlit to refresh when upload happens
        text_key = f"notes_input_v{st.session_state.input_version}"
        notes_text = st.text_area(
            "Raw notes / transcript",
            value=st.session_state.input_text,
            height=400,
            key=text_key,
            placeholder=(
                "Example:\n"
                "Alice: Let's review Q4 goals. Bob, how's the API migration?\n"
                "Bob: About 70% done. Should finish by Friday.\n"
                "Alice: Great. Also we need to ship the pricing page by next Monday.\n"
                "Carol: I'll own that.\n"
                "Decision: freeze scope for v2.0 release.\n"
                "..."
            ),
        )
        # Save what user typed back to session
        st.session_state.input_text = notes_text

        col_a, col_b = st.columns([1, 1])
        with col_a:
            run = st.button("🚀 Generate Summary", type="primary", use_container_width=True)
        with col_b:
            if st.button("🗑️ Clear input", use_container_width=True):
                st.session_state.input_text = ""
                st.session_state.last_upload_name = None
                st.session_state.input_version += 1
                st.rerun()

        if run:
            text_to_use = (notes_text or "").strip()
            if not text_to_use:
                st.error("Please paste or upload some meeting notes.")
            elif len(text_to_use) < 40:
                st.error("Notes are too short to analyze. Paste more content.")
            elif not resolve_api_key():
                st.error("Groq API key is missing. Add GROQ_API_KEY to secrets or environment.")
            else:
                with st.spinner("Analyzing meeting… (this may take 10-30 seconds)"):
                    try:
                        meeting = analyze_meeting(
                            notes=text_to_use,
                            template_key=st.session_state.template,
                            language=st.session_state.language,
                            model=st.session_state.preferred_model,
                        )

                        if st.session_state.enhance_readability:
                            try:
                                enhanced = enhance_readability(text_to_use, st.session_state.preferred_model)
                                meeting["enhanced_notes"] = enhanced
                            except Exception:
                                pass

                        save_meeting(meeting)
                        st.success(f"✅ Analyzed: **{meeting['meeting_title']}**")
                        st.info("Switch to the **Structured Summary** tab to view results.")
                    except Exception as exc:
                        err = str(exc)
                        st.session_state.last_error = err
                        st.error(f"Failed to analyze: {err}")

    # ---------- Tab: Result ----------
    with tab_result:
        meeting = get_active_meeting()
        if not meeting:
            st.info("No meeting analyzed yet. Paste notes in the Input tab and click **Generate Summary**.")
            render_landing()
        else:
            render_meeting(meeting)

    # ---------- Tab: Export ----------
    with tab_export:
        meeting = get_active_meeting()
        if not meeting:
            st.info("Analyze a meeting first, then export from here.")
        else:
            st.subheader(f"Export — {meeting.get('meeting_title','Meeting')}")

            md = build_markdown_report(meeting)
            js = build_json_report(meeting)
            csv = build_actions_csv(meeting)
            pdf_bytes = build_pdf_report(meeting)

            c1, c2 = st.columns(2)
            with c1:
                st.download_button(
                    "📥 Markdown report (.md)",
                    md,
                    file_name=f"meeting_{meeting.get('meeting_date','') or 'summary'}.md",
                    mime="text/markdown",
                    type="primary",
                    use_container_width=True,
                )
            with c2:
                st.download_button(
                    "🧾 Structured data (.json)",
                    js,
                    file_name=f"meeting_{meeting.get('meeting_date','') or 'summary'}.json",
                    mime="application/json",
                    use_container_width=True,
                )

            c3, c4 = st.columns(2)
            with c3:
                st.download_button(
                    "📊 Action items (.csv)",
                    csv,
                    file_name="action_items.csv",
                    mime="text/csv",
                    use_container_width=True,
                )
            with c4:
                if pdf_bytes:
                    st.download_button(
                        "📄 Report (.pdf)",
                        pdf_bytes,
                        file_name=f"meeting_{meeting.get('meeting_date','') or 'summary'}.pdf",
                        mime="application/pdf",
                        use_container_width=True,
                    )
                else:
                    st.button("📄 PDF unavailable (install fpdf2)", disabled=True, use_container_width=True)

            st.markdown("---")
            st.markdown("#### Preview — Markdown")
            with st.expander("Click to expand full markdown", expanded=False):
                st.code(md, language="markdown")

    st.markdown("---")
    st.caption(
        "🔒 **Privacy** — Meeting notes are sent to Groq for AI analysis. "
        "Analyzed meetings are kept in memory for this session only; nothing is persisted to disk."
    )


if __name__ == "__main__":
    main()